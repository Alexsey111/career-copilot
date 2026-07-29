from __future__ import annotations

from app.domain.parse_diagnostics import ParseDiagnosticsReport
from app.services.parse_diagnostics_service import ParseDiagnosticsService
from app.services.resume_parser_service import ParsedResume, ResumeParserService


def _service() -> ParseDiagnosticsService:
    return ParseDiagnosticsService()


def _ru_resume_text() -> str:
    return (
        "Иван Иванов\n"
        "ivan@example.com\n"
        "+7 999 123 45 67\n"
        "ОБО МНЕ\n"
        "Backend-разработчик с опытом 5 лет.\n"
        "ОПЫТ РАБОТЫ\n"
        "Компания А, Senior Python Developer, 2020-2024\n"
        "ОБРАЗОВАНИЕ\n"
        "МГТУ им. Баумана, 2016-2020\n"
        "НАВЫКИ\n"
        "Python, FastAPI, PostgreSQL\n"
    )


def test_block_detection_ru_finds_sections_in_order() -> None:
    report = _service().build_report(
        ParsedResume(text=_ru_resume_text(), detected_format="txt", metadata={"diagnostics_seed": {}})
    )

    assert report.block_order == ["summary", "experience", "education", "skills"]
    skills = next(b for b in report.extracted_blocks if b.kind == "skills")
    assert "Python, FastAPI, PostgreSQL" in skills.sample


def test_block_detection_en_finds_sections() -> None:
    text = (
        "John Doe\njohn@mail.com\n"
        "SUMMARY\nBackend engineer.\n"
        "EXPERIENCE\nAcme, 2018-2024\n"
        "EDUCATION\nMIT, 2014-2018\n"
        "SKILLS\nPython, Go\n"
    )
    report = _service().build_report(
        ParsedResume(text=text, detected_format="txt", metadata={"diagnostics_seed": {}})
    )

    assert report.block_order == ["summary", "experience", "education", "skills"]


def test_prelude_with_contacts_is_not_lost() -> None:
    report = _service().build_report(
        ParsedResume(text=_ru_resume_text(), detected_format="txt", metadata={"diagnostics_seed": {}})
    )
    # Контактный блок (имя + email + телефон) перед первым заголовком не lost.
    assert all("ivan@example.com" not in (b.sample or "") for b in report.lost_blocks)
    assert report.lost_blocks == []


def test_prelude_without_contacts_is_lost() -> None:
    text = (
        "какой-то случайный текст без контактов\n"
        "ещё одна строка\n"
        "ОПЫТ РАБОТЫ\n"
        "Компания А, 2020-2024\n"
    )
    report = _service().build_report(
        ParsedResume(text=text, detected_format="txt", metadata={"diagnostics_seed": {}})
    )

    assert len(report.lost_blocks) == 1
    assert "случайный текст" in report.lost_blocks[0].sample


def test_structural_warnings_long_line() -> None:
    long_line = "x" * 250
    text = f"ОПЫТ РАБОТЫ\n{long_line}\n"
    report = _service().build_report(
        ParsedResume(text=text, detected_format="txt", metadata={"diagnostics_seed": {}})
    )

    codes = [w.code for w in report.structural_warnings]
    assert "long_lines" in codes


def test_structural_warnings_wide_whitespace_and_table() -> None:
    text = "ОПЫТ РАБОТЫ\nКолонка 1      Колонка 2\nЯчейка | Ячейка | Ячейка\n"
    report = _service().build_report(
        ParsedResume(text=text, detected_format="txt", metadata={"diagnostics_seed": {}})
    )

    codes = [w.code for w in report.structural_warnings]
    assert "wide_whitespace_columns" in codes
    assert "table_like" in codes


def test_zero_width_count_from_seed() -> None:
    parser = ResumeParserService()
    text_with_zw = "Python​FastAPI invisible keyword stuffing"
    zw = parser._collect_zero_width(text_with_zw)
    assert zw["count"] == 1

    parsed = ParsedResume(
        text="Python FastAPI invisible keyword stuffing",
        detected_format="txt",
        metadata={"diagnostics_seed": {"zero_width": zw}},
    )
    report = _service().build_report(parsed)

    findings = {f.kind: f for f in report.hidden_text_findings}
    assert "zero_width" in findings
    assert findings["zero_width"].count == 1
    assert findings["zero_width"].severity == "low"


def test_pdf_white_on_white_and_tiny_font_detected() -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    # Белый текст (white-on-white) и tiny-font (size=1) — векторы stuffing.
    page.insert_text((72, 72), "hidden keyword", color=(1, 1, 1), fontsize=12)
    page.insert_text((72, 100), "tiny secret", color=(0, 0, 0), fontsize=1)
    page.insert_text((72, 130), "visible normal text", color=(0, 0, 0), fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()

    # Вызываем _parse_pdf напрямую: autouse-fixture conftest патчит только
    # публичный ``parse``, возвращая фейк без diagnostics_seed.
    parsed = ResumeParserService()._parse_pdf(pdf_bytes)
    report = _service().build_report(parsed)

    kinds = {f.kind for f in report.hidden_text_findings}
    assert "white_on_white" in kinds
    assert "tiny_font" in kinds
    white = next(f for f in report.hidden_text_findings if f.kind == "white_on_white")
    assert white.severity == "high"
    assert "hidden keyword" in white.sample


def test_pdf_metadata_exposure_warning() -> None:
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Some visible resume text", color=(0, 0, 0), fontsize=12)
    doc.set_metadata({"author": "Candidate Name", "producer": "SomeTool", "title": "My Resume"})
    pdf_bytes = doc.tobytes()
    doc.close()

    parsed = ResumeParserService()._parse_pdf(pdf_bytes)
    report = _service().build_report(parsed)

    assert report.metadata_exposure_warning is True
    assert report.file_metadata.author == "Candidate Name"
    assert report.file_metadata.producer == "SomeTool"


def test_docx_vanish_and_white_text_detected() -> None:
    from io import BytesIO

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    doc = Document()
    # Скрытый текст Word (w:vanish).
    p1 = doc.add_paragraph()
    run_vanish = p1.add_run("vanish keyword")
    rpr = run_vanish._element.get_or_add_rPr()
    vanish = OxmlElement("w:vanish")
    rpr.append(vanish)
    # Белый цвет шрифта.
    p2 = doc.add_paragraph()
    run_white = p2.add_run("white on white")
    run_white.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Обычный видимый текст.
    doc.add_paragraph("normal visible experience")

    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parsed = ResumeParserService()._parse_docx(docx_bytes, "x.docx")
    report = _service().build_report(parsed)

    kinds = {f.kind for f in report.hidden_text_findings}
    assert "docx_vanish" in kinds
    assert "white_on_white" in kinds
    vanish_f = next(f for f in report.hidden_text_findings if f.kind == "docx_vanish")
    assert "vanish keyword" in vanish_f.sample


def test_docx_core_properties_exposure() -> None:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("Some visible resume text here")
    cp = doc.core_properties
    cp.author = "Leaked Author"
    cp.last_modified_by = "Editor"

    buf = BytesIO()
    doc.save(buf)

    parsed = ResumeParserService()._parse_docx(buf.getvalue(), "x.docx")
    report = _service().build_report(parsed)

    assert report.metadata_exposure_warning is True
    assert report.file_metadata.author == "Leaked Author"
    assert report.file_metadata.creator_tool == "Editor"


def test_report_as_dict_roundtrip_shape() -> None:
    report = _service().build_report(
        ParsedResume(text=_ru_resume_text(), detected_format="txt", metadata={"diagnostics_seed": {}})
    )
    d = report.as_dict()

    assert d["detected_format"] == "txt"
    assert set(d.keys()) == {
        "detected_format",
        "stats",
        "extracted_blocks",
        "block_order",
        "lost_blocks",
        "structural_warnings",
        "hidden_text_findings",
        "file_metadata",
        "metadata_exposure_warning",
    }
    assert isinstance(d["extracted_blocks"], list) and d["extracted_blocks"]


class _FakePdfPage:
    """Имитация fitz-страницы: ``get_text("dict")`` возвращает структуру
    blocks→lines→spans с ``text``/``color``/``size``. ``color`` — signed ARGB
    int как у fitz (белый = -1 = 0xFFFFFFFF, чёрный = -16777216)."""

    def __init__(self, spans: list[dict]) -> None:
        self._spans = spans

    def get_text(self, mode: str) -> dict:
        line = {"spans": self._spans}
        return {"blocks": [{"lines": [line]}]}


_WHITE = -1  # 0xFFFFFFFF
_BLACK = -16777216  # 0xFF000000
_GRAY = 0xFF666666 & 0xFFFFFFFF  # видимый серый (как у Google Docs артефакта)


def _spans(*triples: tuple[str, int, float]) -> list[dict]:
    return [{"text": t, "color": c, "size": s} for t, c, s in triples]


def test_white_on_white_real_hidden_is_flagged() -> None:
    """Белый текст, которого нигде на странице нет в видимом виде — реальный
    hidden keyword-stuffing, flagged (severity high)."""
    service = ResumeParserService()
    findings: list[dict] = []
    page = _FakePdfPage(_spans(("Python", _BLACK, 12.0), ("SEO spam hidden", _WHITE, 12.0)))
    service._collect_pdf_hidden_spans(page, findings)

    white = [f for f in findings if f["kind"] == "white_on_white"]
    assert len(white) == 1
    assert white[0]["count"] == 1
    assert "hidden" in white[0]["sample"].casefold()


def test_white_on_white_google_docs_artifact_not_flagged() -> None:
    """Ноут #91/#92: Google Docs кладёт поверх видимого текста белую копию того
    же span (render-artifact). Белый дубль видимого текста — НЕ hidden, не
    flagged. Видимый серый «Промпт-инжиниринг» + белая копия → 0 findings."""
    service = ResumeParserService()
    findings: list[dict] = []
    page = _FakePdfPage(
        _spans(
            ("Промпт-инжиниринг", _GRAY, 12.0),
            ("Промпт-инжиниринг", _WHITE, 12.0),
        )
    )
    service._collect_pdf_hidden_spans(page, findings)

    assert not any(f["kind"] == "white_on_white" for f in findings)


def test_white_on_white_artifact_plus_real_hidden() -> None:
    """Смешанный случай: артефакт-дубль (видимая копия есть) игнорируется, а
    настоящий белый hidden (видимой копии нет) — flagged. Счётчик = 1, не 2."""
    service = ResumeParserService()
    findings: list[dict] = []
    page = _FakePdfPage(
        _spans(
            ("Промпт-инжиниринг", _GRAY, 12.0),
            ("Промпт-инжиниринг", _WHITE, 12.0),  # артефакт — игнор
            ("Keyword stuffing", _WHITE, 12.0),  # real hidden — flagged
        )
    )
    service._collect_pdf_hidden_spans(page, findings)

    white = [f for f in findings if f["kind"] == "white_on_white"]
    assert len(white) == 1
    assert white[0]["count"] == 1
    assert "stuffing" in white[0]["sample"].casefold()


def test_tiny_font_still_flagged() -> None:
    """Tiny-font (<2pt) детектится независимо от white-on-white отсева."""
    service = ResumeParserService()
    findings: list[dict] = []
    page = _FakePdfPage(_spans(("Python", _BLACK, 12.0), ("tiny secret", _BLACK, 1.0)))
    service._collect_pdf_hidden_spans(page, findings)

    tiny = [f for f in findings if f["kind"] == "tiny_font"]
    assert len(tiny) == 1
    assert tiny[0]["count"] == 1